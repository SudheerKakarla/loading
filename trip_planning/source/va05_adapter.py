"""Map a saved VA05 workbook into the planner's contract tables.

The adapter reads cached Excel values. It does not recalculate formulas.
A trip rate or case rate of zero is written blank, so that truck or case is
not offered. A ship-to that is missing from Master Customer is still added
from the rate card, with ``in_master`` false. Contract service, when filled,
wins over the master service. Distance is copied and is not used as transit.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd

from baseline_cost import parse_number
from config import TRUCK_TYPES, norm_id
from retrospective import clean_po_numbers

_TRUCK_TOKENS = ("18w", "10w", "6w", "4w")

_VA05_FIELDS = {
    "document": ("sales document", "billing document", "document", "sales doc", "order"),
    "po": ("po number", "po numbers", "purchase order", "customer po", "po"),
    "doc_date": ("document date", "doc date", "created on"),
    "rdd": ("required delivery date", "requested delivery date", "delivery date", "rdd"),
    "ship_to": ("ship to party", "ship to code", "ship to"),
    "sold_to": ("sold to party", "sold to code", "sold to"),
    "cases": ("total ctn convert", "total ctn", "cases"),
    "weight": ("gross weight", "weight kg", "net weight", "weight"),
    "volume": ("volume m3", "volume", "cube"),
}
_MASTER_FIELDS = {
    "ship_to": ("ship to party", "ship to code", "ship to"),
    "sold_to": ("sold to party", "sold to code", "sold to"),
    "province": ("province",),
    "group": ("customer group", "group"),
    "service": ("service type", "service"),
    "sla": ("sla days", "lead time", "sla"),
    "name": ("customer name", "name"),
    "segment": ("segment",),
    "order_type": ("order type",),
}
_RATE_PLAIN = {
    "ship_to": ("ship to party", "ship to code", "ship to"),
    "sold_to": ("sold to party", "sold to"),
    "province": ("province",),
    "group": ("customer group", "group"),
    "service": ("service type", "contract service", "service"),
    "name": ("customer name", "name"),
    "segment": ("segment",),
    "distance_km": ("round trip distance", "distance km", "distance"),
    "remark": ("multi drop", "multidrop", "remarks", "remark", "note"),
    "ltl_mt": ("ltl mt", "case mt", "modern trade", "mt case"),
    "ltl_gt": ("ltl gt", "case gt", "general trade", "gt case"),
}


@dataclass(frozen=True)
class WorkbookExtract:
    """Contract tables plus source notes. Diagnostics are counts, not prices."""

    orders: pd.DataFrame
    customers: pd.DataFrame
    rates: pd.DataFrame
    trucks: pd.DataFrame
    issues: pd.DataFrame
    diagnostics: dict[str, int]


def read_workbook(path: Path) -> WorkbookExtract:
    """Read ``path`` and return validated orders, customers, rates, and trucks."""
    workbook = Path(path)
    if not workbook.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook}.")
    _openpyxl()
    va05 = _sheet_matrix(workbook, "VA05")
    master = _sheet_matrix(workbook, "Master Customer")
    rate_card = _sheet_matrix(workbook, "TPT Rate Card")
    pivot = _sheet_matrix(workbook, "Pivot_Summary")
    lines, line_count = _va05_lines(va05)
    orders = _aggregate_orders(lines)
    master_rows = _master_rows(master)
    rate_rows, rate_notes = _rate_rows(rate_card)
    customers, customer_notes = _customers(master_rows, rate_rows)
    trucks, truck_notes = _trucks(pivot)
    issues = _issues(customer_notes + rate_notes + truck_notes)
    diagnostics = {
        "data_order_lines": line_count,
        "data_orders": int(len(orders)),
        "data_customers": int(len(customers)),
        "data_customers_not_in_master": int((customers["in_master"] == "false").sum()) if len(customers) else 0,
        "data_rates": int(len(rate_rows)),
        "data_service_conflicts": int((customers["service_conflict"] == "true").sum()) if len(customers) else 0,
        "data_zero_trip_rates_blanked": rate_notes_count(rate_notes, "zero_trip_rate_blanked"),
        "data_zero_case_rates_blanked": rate_notes_count(rate_notes, "zero_case_rate_blanked"),
    }
    return WorkbookExtract(orders, customers, _rates_frame(rate_rows), trucks, issues, diagnostics)


def rate_notes_count(notes: list[dict], issue: str) -> int:
    """Count adapter notes of one kind. Public so the builder summary can reuse it."""
    return sum(1 for note in notes if note["issue"] == issue)


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read a VA05 workbook.") from exc


def _sheet_matrix(path: Path, sheet_name: str) -> list[list]:
    import openpyxl

    raw_book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    cached_book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in raw_book.sheetnames or sheet_name not in cached_book.sheetnames:
            raise ValueError(f"Workbook is missing sheet {sheet_name}. Found: {', '.join(raw_book.sheetnames)}.")
        raw_rows = raw_book[sheet_name].iter_rows(values_only=True)
        cached_rows = cached_book[sheet_name].iter_rows(values_only=True)
        matrix = []
        for row_number, (raw_row, cached_row) in enumerate(zip(raw_rows, cached_rows), start=1):
            width = max(len(raw_row), len(cached_row))
            values = []
            for column in range(width):
                raw_value = raw_row[column] if column < len(raw_row) else None
                cached_value = cached_row[column] if column < len(cached_row) else None
                if isinstance(raw_value, str) and raw_value.startswith("="):
                    if cached_value is None:
                        raise ValueError(
                            f"{sheet_name} row {row_number} has a formula with no cached value."
                        )
                    values.append(cached_value)
                else:
                    values.append(raw_value)
            matrix.append(values)
        return matrix
    finally:
        raw_book.close()
        cached_book.close()


def _header_key(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("³", "3").replace("²", "2").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _locate_header(matrix: list[list], fields: dict[str, tuple[str, ...]], required: tuple[str, ...], sheet: str) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(matrix[:20]):
        found = _map_header(row, fields, sheet)
        if all(name in found for name in required):
            return index, found
    raise ValueError(f"{sheet} has no header row with {', '.join(required)}.")


def _map_header(row: list, fields: dict[str, tuple[str, ...]], sheet: str) -> dict[str, int]:
    found: dict[str, int] = {}
    for column, cell in enumerate(row):
        key = _header_key(cell)
        if not key:
            continue
        for field_name, aliases in fields.items():
            if key in aliases:
                if field_name in found:
                    raise ValueError(f"{sheet} has two columns for {field_name}.")
                found[field_name] = column
                break
    return found


def _cell(row: list, columns: dict[str, int], name: str):
    index = columns.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


def _blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and value != value:
        return True
    return str(value).strip().lower() in {"", "nan", "none", "<na>", "nat"}


def _require_id(value: object, where: str) -> str:
    if isinstance(value, bool):
        raise ValueError(f"{where} is not an id.")
    identity = norm_id(value)
    if not identity:
        raise ValueError(f"{where} is missing.")
    return identity


def _require_quantity(value: object, where: str) -> float:
    if isinstance(value, bool) or (isinstance(value, str) and value.strip().lower() in {"true", "false"}):
        raise ValueError(f"{where} must be a finite number, not a boolean.")
    try:
        number = parse_number(value)
    except ValueError as exc:
        raise ValueError(f"{where} must be a finite number.") from exc
    if number is None:
        raise ValueError(f"{where} is missing.")
    if number < 0:
        raise ValueError(f"{where} is negative.")
    return float(number)


def _optional_money(value: object, where: str) -> float | None:
    if isinstance(value, bool):
        raise ValueError(f"{where} must be a finite number, not a boolean.")
    try:
        number = parse_number(value)
    except ValueError as exc:
        raise ValueError(f"{where} must be a finite number.") from exc
    if number is None:
        return None
    if number < 0:
        raise ValueError(f"{where} is negative.")
    return float(number)


def _excel_date(value: object, where: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise ValueError(f"{where} is not a date.")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value != value or value in {float("inf"), float("-inf")}:
            raise ValueError(f"{where} is not a date.")
        from openpyxl.utils.datetime import from_excel

        try:
            converted = from_excel(value)
        except Exception as exc:
            raise ValueError(f"{where} is not a date.") from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
    text = "" if value is None else str(value).strip()
    if len(text) >= 10 and text[4:5] == "-" and text[7:8] == "-":
        try:
            return date.fromisoformat(text[:10])
        except ValueError as exc:
            raise ValueError(f"{where} is not a date.") from exc
    raise ValueError(f"{where} is not a date.")


def _text(value: object) -> str:
    if _blank(value):
        return ""
    return str(value).strip()


def _service(value: object) -> str:
    return _text(value).upper()


def _sla_days(value: object, ship_to: str) -> str:
    if _blank(value):
        return ""
    text = _text(value).upper().replace(" ", "")
    if text.startswith("D+"):
        text = text[2:]
    elif text.endswith("D") and text[:-1].isdigit():
        text = text[:-1]
    number = _optional_money(text, f"Master Customer ship-to {ship_to} sla")
    if number is None or not float(number).is_integer() or number < 0:
        raise ValueError(f"Master Customer ship-to {ship_to} has an invalid SLA {value!r}.")
    return str(int(number))


def _va05_lines(matrix: list[list]) -> tuple[list[dict], int]:
    header_index, columns = _locate_header(
        matrix,
        _VA05_FIELDS,
        ("document", "ship_to", "sold_to", "doc_date", "rdd", "cases", "weight", "volume"),
        "VA05",
    )
    lines = []
    for offset, row in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        if all(_blank(cell) for cell in row):
            continue
        where = f"VA05 row {offset}"
        lines.append(
            {
                "document": _require_id(_cell(row, columns, "document"), f"{where} document"),
                "sold_to": _require_id(_cell(row, columns, "sold_to"), f"{where} sold-to"),
                "ship_to": _require_id(_cell(row, columns, "ship_to"), f"{where} ship-to"),
                "doc_date": _excel_date(_cell(row, columns, "doc_date"), f"{where} document date"),
                "rdd": _excel_date(_cell(row, columns, "rdd"), f"{where} delivery date"),
                "cases": _require_quantity(_cell(row, columns, "cases"), f"{where} cases"),
                "weight_kg": _require_quantity(_cell(row, columns, "weight"), f"{where} weight"),
                "volume_m3": _require_quantity(_cell(row, columns, "volume"), f"{where} volume"),
                "po": _text(_cell(row, columns, "po")),
            }
        )
    return lines, len(lines)


def _aggregate_orders(lines: list[dict]) -> pd.DataFrame:
    grouped: dict[tuple[str, date], list[dict]] = {}
    for line in lines:
        grouped.setdefault((line["document"], line["rdd"]), []).append(line)
    rows = []
    for (document, rdd), members in grouped.items():
        ship_tos = {member["ship_to"] for member in members}
        sold_tos = {member["sold_to"] for member in members}
        label = f"{document}@{rdd.isoformat()}"
        if len(ship_tos) != 1 or len(sold_tos) != 1:
            raise ValueError(f"Order {label} has conflicting ship-to or sold-to values.")
        rows.append(
            {
                "order_id": label,
                "ship_to": next(iter(ship_tos)),
                "doc_date": max(member["doc_date"] for member in members).isoformat(),
                "rdd": rdd.isoformat(),
                "cases": sum(member["cases"] for member in members),
                "weight_kg": sum(member["weight_kg"] for member in members),
                "volume_m3": sum(member["volume_m3"] for member in members),
                "po_numbers": clean_po_numbers(";".join(member["po"] for member in members if member["po"])),
            }
        )
    rows.sort(key=lambda row: row["order_id"])
    return pd.DataFrame(
        rows,
        columns=["order_id", "ship_to", "doc_date", "rdd", "cases", "weight_kg", "volume_m3", "po_numbers"],
    )


def _table_rows(matrix: list[list], fields: dict[str, tuple[str, ...]], required: tuple[str, ...], sheet: str) -> tuple[int, dict[str, int], list[list]]:
    header_index, columns = _locate_header(matrix, fields, required, sheet)
    body = []
    for row in matrix[header_index + 1 :]:
        if all(_blank(cell) for cell in row):
            continue
        body.append(row)
    return header_index, columns, body


def _master_rows(matrix: list[list]) -> list[dict]:
    _header, columns, body = _table_rows(matrix, _MASTER_FIELDS, ("ship_to",), "Master Customer")
    rows = []
    seen: set[str] = set()
    for offset, row in enumerate(body, start=2):
        ship_to = _require_id(_cell(row, columns, "ship_to"), f"Master Customer row {offset} ship-to")
        if ship_to in seen:
            raise ValueError(f"Master Customer ship-to {ship_to} is duplicated.")
        seen.add(ship_to)
        rows.append(
            {
                "ship_to": ship_to,
                "sold_to": "" if _blank(_cell(row, columns, "sold_to")) else norm_id(_cell(row, columns, "sold_to")),
                "province": _text(_cell(row, columns, "province")),
                "group": _text(_cell(row, columns, "group")),
                "service": _service(_cell(row, columns, "service")),
                "sla_days": _sla_days(_cell(row, columns, "sla"), ship_to),
                "name": _text(_cell(row, columns, "name")),
                "segment": _text(_cell(row, columns, "segment")),
                "order_type": _text(_cell(row, columns, "order_type")),
            }
        )
    return rows


def _rate_field(header: str) -> str | None:
    if header in {alias for aliases in _RATE_PLAIN.values() for alias in aliases}:
        for field_name, aliases in _RATE_PLAIN.items():
            if header in aliases:
                return field_name
    tokens = set(header.split())
    for truck in _TRUCK_TOKENS:
        if truck in tokens or header.endswith(truck):
            if "handl" in header or "unload" in header:
                return f"handling_{truck}"
            if any(word in header for word in ("trip", "rate", "price", "baht", "freight")) or header == truck:
                return f"trip_{truck}"
    return None


def _rate_columns(row: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for column, cell in enumerate(row):
        key = _header_key(cell)
        if not key:
            continue
        field_name = _rate_field(key)
        if field_name is None:
            continue
        if field_name in found:
            raise ValueError(f"TPT Rate Card has two columns for {field_name}.")
        found[field_name] = column
    if "ship_to" not in found:
        return {}
    return found


def _rate_rows(matrix: list[list]) -> tuple[list[dict], list[dict]]:
    columns: dict[str, int] = {}
    header_index = None
    for index, row in enumerate(matrix[:25]):
        mapped = _rate_columns(row)
        if mapped:
            columns = mapped
            header_index = index
            break
    if header_index is None:
        raise ValueError("TPT Rate Card has no ship-to header.")
    rows = []
    notes = []
    seen: set[str] = set()
    for offset, row in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        if all(_blank(cell) for cell in row):
            continue
        ship_to = _require_id(_cell(row, columns, "ship_to"), f"TPT Rate Card row {offset} ship-to")
        if ship_to in seen:
            raise ValueError(f"TPT Rate Card ship-to {ship_to} is duplicated.")
        seen.add(ship_to)
        record = {
            "ship_to": ship_to,
            "sold_to": "" if _blank(_cell(row, columns, "sold_to")) else norm_id(_cell(row, columns, "sold_to")),
            "province": _text(_cell(row, columns, "province")),
            "group": _text(_cell(row, columns, "group")),
            "service": _service(_cell(row, columns, "service")),
            "name": _text(_cell(row, columns, "name")),
            "segment": _text(_cell(row, columns, "segment")),
            "distance_km": _distance(_cell(row, columns, "distance_km"), ship_to),
            "multidrop": _remark(_cell(row, columns, "remark")),
        }
        for truck in ("4w", "6w", "10w", "18w"):
            record[f"trip_{truck}"] = _blank_zero_rate(
                _cell(row, columns, f"trip_{truck}"),
                f"TPT Rate Card {ship_to} trip {truck.upper()}",
                ship_to,
                "zero_trip_rate_blanked",
                notes,
            )
            handling = _optional_money(
                _cell(row, columns, f"handling_{truck}"),
                f"TPT Rate Card {ship_to} handling {truck.upper()}",
            )
            record[f"handling_{truck}"] = "" if handling is None else handling
        for column in ("ltl_mt", "ltl_gt"):
            record[column] = _blank_zero_rate(
                _cell(row, columns, column),
                f"TPT Rate Card {ship_to} {column}",
                ship_to,
                "zero_case_rate_blanked",
                notes,
            )
        rows.append(record)
    return rows, notes


def _blank_zero_rate(value: object, where: str, ship_to: str, issue: str, notes: list[dict]):
    number = _optional_money(value, where)
    if number is None:
        return ""
    if abs(number) <= 1e-9:
        notes.append({"issue": issue, "key": ship_to, "detail": where})
        return ""
    return number


def _distance(value: object, ship_to: str):
    number = _optional_money(value, f"TPT Rate Card {ship_to} distance")
    return "" if number is None else number


def _remark(value: object) -> str:
    text = _text(value)
    if text.lower() in {"", "0", "false", "no", "n"}:
        return ""
    return text


def _customers(master_rows: list[dict], rate_rows: list[dict]) -> tuple[pd.DataFrame, list[dict]]:
    by_rate = {row["ship_to"]: row for row in rate_rows}
    notes = []
    rows = []
    seen: set[str] = set()
    for master in master_rows:
        ship_to = master["ship_to"]
        seen.add(ship_to)
        contract = by_rate.get(ship_to, {}).get("service", "")
        master_service = master["service"]
        effective = contract or master_service
        conflict = bool(master_service and contract and master_service != contract)
        if conflict:
            notes.append(
                {
                    "issue": "service_conflict",
                    "key": ship_to,
                    "detail": f"master {master_service} vs contract {contract}",
                }
            )
        rows.append(
            {
                "ship_to": ship_to,
                "sold_to": master["sold_to"],
                "province": master["province"] or by_rate.get(ship_to, {}).get("province", ""),
                "group": master["group"] or by_rate.get(ship_to, {}).get("group", ""),
                "service": effective,
                "sla_days": master["sla_days"],
                "name": master["name"] or by_rate.get(ship_to, {}).get("name", ""),
                "segment": master["segment"] or by_rate.get(ship_to, {}).get("segment", ""),
                "order_type": master["order_type"],
                "service_master": master_service,
                "service_contract": contract,
                "service_conflict": "true" if conflict else "false",
                "in_master": "true",
            }
        )
    for rate in rate_rows:
        if rate["ship_to"] in seen:
            continue
        notes.append(
            {
                "issue": "not_in_master_customer",
                "key": rate["ship_to"],
                "detail": "built from the rate card",
            }
        )
        rows.append(
            {
                "ship_to": rate["ship_to"],
                "sold_to": rate["sold_to"],
                "province": rate["province"],
                "group": rate["group"],
                "service": rate["service"],
                "sla_days": "",
                "name": rate["name"],
                "segment": rate["segment"],
                "order_type": "",
                "service_master": "",
                "service_contract": rate["service"],
                "service_conflict": "false",
                "in_master": "false",
            }
        )
    rows.sort(key=lambda row: row["ship_to"])
    columns = [
        "ship_to",
        "sold_to",
        "province",
        "group",
        "service",
        "sla_days",
        "name",
        "segment",
        "order_type",
        "service_master",
        "service_contract",
        "service_conflict",
        "in_master",
    ]
    return pd.DataFrame(rows, columns=columns), notes


def _rates_frame(rows: list[dict]) -> pd.DataFrame:
    columns = [
        "ship_to",
        "province",
        "group",
        "service",
        "trip_4w",
        "trip_6w",
        "trip_10w",
        "trip_18w",
        "handling_4w",
        "handling_6w",
        "handling_10w",
        "handling_18w",
        "ltl_mt",
        "ltl_gt",
        "multidrop",
        "distance_km",
        "segment",
    ]
    ordered = sorted(rows, key=lambda row: row["ship_to"])
    return pd.DataFrame(ordered, columns=columns)


def _trucks(matrix: list[list]) -> tuple[pd.DataFrame, list[dict]]:
    header = None
    columns: dict[str, int] = {}
    for index, row in enumerate(matrix):
        labels = {_header_key(cell): column for column, cell in enumerate(row) if _header_key(cell)}
        type_at = next((labels[name] for name in ("truck type", "type") if name in labels), None)
        kg_at = next((labels[name] for name in ("capacity kg", "kg") if name in labels), None)
        m3_at = next((labels[name] for name in ("capacity m3", "m3", "cbm") if name in labels), None)
        if type_at is None or kg_at is None or m3_at is None:
            continue
        header = index
        columns = {"truck_type": type_at, "capacity_kg": kg_at, "capacity_m3": m3_at}
        break
    if header is None:
        raise ValueError("Pivot_Summary has no Type / KG / M3 truck table.")
    found: dict[str, tuple[object, object]] = {}
    notes = []
    for offset, row in enumerate(matrix[header + 1 :], start=header + 2):
        raw_type = _cell(row, columns, "truck_type")
        if _blank(raw_type):
            if found:
                break
            continue
        truck_type = _text(raw_type).upper()
        if truck_type not in TRUCK_TYPES:
            notes.append({"issue": "truck_type_ignored", "key": truck_type, "detail": "not a 4W/6W/10W/18W capacity"})
            continue
        if truck_type in found:
            raise ValueError(f"Pivot_Summary lists {truck_type} more than once.")
        kg = _capacity(_cell(row, columns, "capacity_kg"), f"Pivot_Summary {truck_type} kg")
        m3 = _capacity(_cell(row, columns, "capacity_m3"), f"Pivot_Summary {truck_type} m3")
        found[truck_type] = (kg, m3)
    if "18W" not in found:
        found["18W"] = ("", "")
        notes.append({"issue": "truck_without_capacity", "key": "18W", "detail": "no 18W row in Pivot_Summary"})
    for truck_type, (kg, m3) in found.items():
        if kg == "" or m3 == "":
            if truck_type == "18W" and any(note["key"] == "18W" for note in notes):
                continue
            notes.append({"issue": "truck_without_capacity", "key": truck_type, "detail": "capacity is blank"})
    rows = [
        {"truck_type": truck_type, "capacity_kg": found[truck_type][0], "capacity_m3": found[truck_type][1]}
        for truck_type in TRUCK_TYPES
        if truck_type in found
    ]
    return pd.DataFrame(rows, columns=["truck_type", "capacity_kg", "capacity_m3"]), notes


def _capacity(value: object, where: str):
    number = _optional_money(value, where)
    return "" if number is None else number


def _issues(notes: list[dict]) -> pd.DataFrame:
    rows = sorted(notes, key=lambda note: (note["issue"], str(note["key"])))
    return pd.DataFrame(rows, columns=["issue", "key", "detail"])
